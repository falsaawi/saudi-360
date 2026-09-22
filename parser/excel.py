"""
Sheet -> tidy rows.

GASTAT workbooks are laid out for reading, not loading, and the same logical
table appears under English names, Arabic names and numeric codes across
releases. So nothing here keys on sheet names. Each sheet is reduced to a
*fingerprint* describing its shape, and a fingerprint selects a *recipe* that
knows how to read that shape.

A sheet whose shape no recipe claims is not guessed at — it goes to the reject
queue with a reason, so a layout change surfaces as an unparsed sheet rather
than as silently wrong numbers.
"""

from __future__ import annotations

import datetime as dt
import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Callable, Iterable

Grid = list[list[Any]]

# Bounds on the scan. The CPI category sheet alone runs to 763 rows, so these
# are set well above any table seen rather than at a round guess -- a cap below
# the table length truncates it silently and drops whole categories.
MAX_SCAN_ROWS = 5000
MAX_SCAN_COLS = 80


# --------------------------------------------------------------------------
# loading
# --------------------------------------------------------------------------

def load_grid(path: Path, sheet: int | str = 0) -> tuple[str, Grid]:
    """Read one sheet as a plain grid. Handles both .xlsx and legacy .xls."""
    suffix = path.suffix.lower()
    if suffix == ".xlsx":
        import openpyxl

        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        try:
            ws = wb[wb.sheetnames[sheet]] if isinstance(sheet, int) else wb[sheet]
            grid = [
                list(row)
                for row in ws.iter_rows(
                    max_row=MAX_SCAN_ROWS, max_col=MAX_SCAN_COLS, values_only=True
                )
            ]
            return ws.title, grid
        finally:
            wb.close()

    if suffix == ".xls":
        import xlrd

        book = xlrd.open_workbook(path, formatting_info=False)
        try:
            ws = book.sheet_by_index(sheet) if isinstance(sheet, int) else book.sheet_by_name(sheet)
            grid: Grid = []
            for r in range(min(ws.nrows, MAX_SCAN_ROWS)):
                row = []
                for c in range(min(ws.ncols, MAX_SCAN_COLS)):
                    cell = ws.cell(r, c)
                    value = cell.value
                    if cell.ctype == xlrd.XL_CELL_DATE:
                        value = dt.datetime(*xlrd.xldate_as_tuple(value, book.datemode))
                    elif cell.ctype == xlrd.XL_CELL_EMPTY:
                        value = None
                    row.append(value)
                grid.append(row)
            return ws.name, grid
        finally:
            book.release_resources()

    raise ValueError(f"unsupported workbook type: {path.suffix}")


def sheet_names(path: Path) -> list[str]:
    if path.suffix.lower() == ".xlsx":
        import openpyxl

        wb = openpyxl.load_workbook(path, read_only=True)
        try:
            return list(wb.sheetnames)
        finally:
            wb.close()
    import xlrd

    book = xlrd.open_workbook(path, formatting_info=False)
    try:
        return list(book.sheet_names())
    finally:
        book.release_resources()


# --------------------------------------------------------------------------
# cell helpers
# --------------------------------------------------------------------------

def text(cell: Any) -> str:
    """A cell as comparable text. Bilingual cells keep both halves."""
    if cell is None:
        return ""
    if isinstance(cell, (dt.datetime, dt.date)):
        return cell.strftime("%Y-%m-%d")
    return " ".join(str(cell).split())


def is_number(cell: Any) -> bool:
    if isinstance(cell, bool) or cell is None:
        return False
    if isinstance(cell, (int, float)):
        return True
    try:
        float(str(cell).replace(",", "").strip())
        return True
    except ValueError:
        return False


def as_number(cell: Any) -> float | None:
    if isinstance(cell, bool) or cell is None:
        return None
    if isinstance(cell, (int, float)):
        return float(cell)
    try:
        return float(str(cell).replace(",", "").strip())
    except ValueError:
        return None


ARABIC_CHAR = re.compile(r"[؀-ۿﭐ-﻿]")
LATIN_CHAR = re.compile(r"[A-Za-z]")


def split_bilingual(cell: Any) -> tuple[str, str]:
    """Separate the English and Arabic halves of a label.

    Returns (english, arabic); either may be empty.

    The two languages arrive together in several shapes: on separate lines in one
    cell, separated by a pipe, or simply run together with a space when a header
    is stacked over several rows ("الأرقام القياسية Index Numbers الرياض Riyadh").
    Splitting only on newlines and pipes sent that last form entirely to Arabic
    and left the English label empty -- which is what made whole tables' columns
    collapse into one unidentifiable series. So the split is per word, by script,
    with digits and punctuation staying with the run they interrupt.
    """
    raw = str(cell or "")
    if not raw.strip():
        return "", ""

    arabic: list[str] = []
    latin: list[str] = []
    current: list[str] | None = None
    for token in re.split(r"[\s\n|]+", raw):
        if not token:
            continue
        is_ar = bool(ARABIC_CHAR.search(token))
        is_en = bool(LATIN_CHAR.search(token))
        if is_ar and is_en:
            # One token carrying both scripts: split it at the boundary.
            ar_part = "".join(ch for ch in token if not LATIN_CHAR.match(ch))
            en_part = "".join(ch for ch in token if not ARABIC_CHAR.match(ch))
            if ar_part.strip():
                arabic.append(ar_part.strip())
            if en_part.strip():
                latin.append(en_part.strip())
            current = None
            continue
        if is_ar:
            current = arabic
        elif is_en:
            current = latin
        elif current is None:
            # Leading digits or punctuation, before either script has appeared.
            current = latin
        current.append(token)

    def clean(parts: list[str]) -> str:
        out = " ".join(parts).strip(" -–:,|")
        return " ".join(out.split())

    return clean(latin), clean(arabic)


# --------------------------------------------------------------------------
# fingerprinting
# --------------------------------------------------------------------------

@dataclass(frozen=True)
class Fingerprint:
    """A sheet's shape, independent of its name and its numbers."""

    n_sheets: int
    header_row: int | None      # the row naming the level column, if found
    header_start: int | None    # first row of the whole header block
    first_data_row: int | None
    label_col: int | None
    n_numeric_cols: int
    has_level_col: bool

    def key(self) -> str:
        return (
            f"s{self.n_sheets}/h{self.header_start}-{self.first_data_row}"
            f"/l{self.label_col}/n{self.n_numeric_cols}/{int(self.has_level_col)}"
        )


LEVEL_TOKENS = ("level", "الدرجة")


def fingerprint(grid: Grid, n_sheets: int) -> Fingerprint:
    header_row = None
    for r, row in enumerate(grid[:20]):
        joined = " ".join(text(c).lower() for c in row[:4])
        if any(tok in joined for tok in LEVEL_TOKENS):
            header_row = r
            break

    first_data_row = None
    for r, row in enumerate(grid):
        if header_row is not None and r <= header_row:
            continue
        # a data row starts with a small integer level
        if row and is_number(row[0]) and float(as_number(row[0]) or -1).is_integer():
            if 0 <= int(as_number(row[0])) <= 9:
                first_data_row = r
                break

    # The label column is the one carrying the most distinct names down the
    # block -- not simply the first text column. Several layouts put a near
    # constant classification code ("-") to the left of the real labels, and
    # taking the first text column silently labels every row "-".
    label_col = None
    if first_data_row is not None:
        best = 0
        for c in range(1, min(8, max(len(r) for r in grid[first_data_row:first_data_row + 40] or [[]]) or 1)):
            values = set()
            for row in grid[first_data_row : first_data_row + 60]:
                if c < len(row) and text(row[c]) and not is_number(row[c]):
                    values.add(text(row[c]))
            if len(values) > best:
                best, label_col = len(values), c
        if best < 2:
            label_col = None

    n_numeric = 0
    if first_data_row is not None:
        n_numeric = sum(1 for c in grid[first_data_row] if is_number(c))

    # The header block is every row above the data that still looks like header:
    # two or more filled cells, blanks tolerated. It is found by walking up from
    # the data rather than from the row naming the level column, because some
    # layouts put that row LAST and the group headings ("Index Numbers",
    # "Percent Change") sit above it.
    header_start = None
    if first_data_row is not None:
        header_start = first_data_row
        for r in range(first_data_row - 1, max(-1, first_data_row - 15), -1):
            cells = [text(c) for c in grid[r] if text(c)]
            # A title banner ends the header: one lone cell of prose. A lone
            # short cell does not -- several layouts put "Level" on a row by
            # itself directly above the data.
            if len(cells) == 1 and len(cells[0]) > 30:
                break
            header_start = r
        while header_start < first_data_row and not any(text(c) for c in grid[header_start]):
            header_start += 1        # trim blank rows off the top of the block

    return Fingerprint(
        n_sheets=n_sheets,
        header_row=header_row,
        header_start=header_start,
        first_data_row=first_data_row,
        label_col=label_col,
        n_numeric_cols=n_numeric,
        has_level_col=header_row is not None,
    )


# --------------------------------------------------------------------------
# recipes
# --------------------------------------------------------------------------

@dataclass
class Observation:
    level: int
    label_en: str
    label_ar: str
    weight: float | None
    value: float
    column: int


@dataclass
class Recipe:
    """How to read one shape of sheet."""

    name: str
    claims: Callable[[Fingerprint], bool]
    read: Callable[..., list[Observation]]
    note: str = ""


class Rejected(Exception):
    """This sheet is not something the recipe can honestly read."""


def _weight_column(grid: Grid, fp: Fingerprint) -> int | None:
    """The weight column, by header text above the data."""
    if fp.header_start is None or fp.first_data_row is None:
        return None
    for c in range((fp.label_col or 0) + 1, MAX_SCAN_COLS):
        above = " ".join(
            normalize(text(grid[r][c]))
            for r in range(fp.header_start, fp.first_data_row)
            if c < len(grid[r])
        )
        if "weight" in above or "وزن" in above or "الاوزان" in above or "الاهميه" in above:
            return c
    return None


MONTH_WORDS = {
    "january": 1, "jan": 1, "يناير": 1,
    "february": 2, "feb": 2, "فبراير": 2,
    "march": 3, "mar": 3, "مارس": 3,
    "april": 4, "apr": 4, "أبريل": 4, "ابريل": 4,
    "may": 5, "مايو": 5,
    "june": 6, "jun": 6, "يونيو": 6,
    "july": 7, "jul": 7, "يوليو": 7,
    "august": 8, "aug": 8, "أغسطس": 8, "اغسطس": 8,
    "september": 9, "sep": 9, "sept": 9, "سبتمبر": 9,
    "october": 10, "oct": 10, "أكتوبر": 10, "اكتوبر": 10,
    "november": 11, "nov": 11, "نوفمبر": 11,
    "december": 12, "dec": 12, "ديسمبر": 12,
}

# Matched against normalize()d header text, so these are written folded:
# no tatweel, no diacritics, alef and ya variants collapsed.
INDEX_WORDS = ("index", "الارقام القياسيه", "الرقم القياسي", "القياسيه", "القياسي")
CHANGE_WORDS = ("change", "تغير", "معدل", "نسبه التغير")


TATWEEL = "ـ"
DIACRITICS = "".join(chr(c) for c in range(0x064B, 0x0653))
ARABIC_FOLD = {
    "آ": "ا", "أ": "ا", "إ": "ا",  # alef variants
    "ى": "ي",                                          # alef maqsura
    "ة": "ه",                                          # ta marbuta
}


def normalize(value: str) -> str:
    """Fold a header string so Arabic matches regardless of presentation.

    GASTAT stretches Arabic headers with tatweel for justification, so the same
    words appear as both "الأرقام القياسية" and "الأرقـــام القياســـية". Without
    folding, a keyword match silently misses and every index column looks like a
    percent-change column.
    """
    out = value.lower().replace(TATWEEL, "")
    out = "".join(ch for ch in out if ch not in DIACRITICS)
    return "".join(ARABIC_FOLD.get(ch, ch) for ch in out)


def _effective(grid: Grid, r: int, col: int) -> tuple[str, bool]:
    """A header cell's value, inheriting from the left when it spans columns.

    Merged header cells read as a value in the first column and blanks after it,
    so a column's own heading is often several cells to the left.
    Returns (value, owned) where owned is False if the value was inherited.
    """
    row = grid[r]
    if col < len(row) and text(row[col]):
        return text(row[col]), True
    for c in range(min(col, len(row)) - 1, -1, -1):
        if text(row[c]):
            return text(row[c]), False
    return "", False


def _group_label(grid: Grid, fp: Fingerprint, col: int) -> str:
    """Every heading a column sits under, normalized and joined.

    The whole header block is scanned, not one row: these sheets routinely put
    the Arabic heading on one row and its English equivalent on the next, and
    either one is enough to classify the column.
    """
    if fp.header_start is None or fp.first_data_row is None:
        return ""
    labels = [
        normalize(_effective(grid, r, col)[0])
        for r in range(fp.header_start, fp.first_data_row)
    ]
    return " | ".join(v for v in labels if v)


def column_period(grid: Grid, fp: Fingerprint, col: int) -> tuple[int, int] | None:
    """The (year, month) a column reports, read from the header rows above it.

    month 0 means the column is annual. Returns None when the header does not
    say, which is treated as a reject rather than a guess.
    """
    if fp.header_start is None or fp.first_data_row is None:
        return None
    year = month = None
    for r in range(fp.header_start, fp.first_data_row):
        if col < len(grid[r]) and isinstance(grid[r][col], (dt.datetime, dt.date)):
            cell = grid[r][col]
            return cell.year, cell.month

        # A year written once above a group of month columns applies to all of
        # them, so an inherited value counts for the year but never for the
        # month -- otherwise every column would take its neighbour's month.
        value, owned = _effective(grid, r, col)
        token = normalize(value)
        if not token:
            continue
        if owned:
            for word, num in MONTH_WORDS.items():
                if re.search(r"(?<![^\W\d_])" + re.escape(normalize(word)), token):
                    month = num
                    break
        found_year = re.search(r"(?<!\d)(19|20)\d{2}(?!\d)", token)
        if found_year:
            year = int(found_year.group(0))
    if year is None:
        return None
    return year, month or 0


def index_base(grid: Grid) -> str | None:
    """The index base year a sheet states, e.g. "2018=100".

    GASTAT rebased the CPI from 2018=100 to 2023=100 partway through the series.
    Values on different bases are not comparable, so the base travels with every
    observation: splicing across it silently produces a fictitious price fall.
    """
    for row in grid[:16]:
        for cell in row:
            found = re.search(r"(20\d\d)\s*=\s*100", text(cell))
            if found:
                return found.group(1)
    return None


def _is_arabic(value: str) -> bool:
    return bool(re.search(r"[؀-ۿ]", value))


def _other_language_column(grid: Grid, fp: Fingerprint) -> int | None:
    """A second label column holding the other language, if there is one.

    Some releases put Arabic and English in one cell; others give each its own
    column. Without this the English label comes back empty on every row of the
    bilingual-column layouts.
    """
    if fp.label_col is None or fp.first_data_row is None:
        return None
    block = grid[fp.first_data_row : fp.first_data_row + 60]
    primary_arabic = sum(
        _is_arabic(text(row[fp.label_col]))
        for row in block
        if fp.label_col < len(row) and text(row[fp.label_col])
    )
    primary_is_arabic = primary_arabic > len(block) / 4

    best, best_col = 0, None
    for c in range(1, MAX_SCAN_COLS):
        if c == fp.label_col:
            continue
        values = {
            text(row[c])
            for row in block
            if c < len(row) and text(row[c]) and not is_number(row[c])
        }
        if len(values) < 5:
            continue
        arabic = sum(_is_arabic(v) for v in values)
        this_is_arabic = arabic > len(values) / 2
        if this_is_arabic != primary_is_arabic and len(values) > best:
            best, best_col = len(values), c
    return best_col


def read_indicator_table(
    grid: Grid, fp: Fingerprint, want: tuple[int, int] | None = None
) -> list[Observation]:
    """A level/label/weight table whose remaining columns are periods.

    The column to read is the one whose own header says it holds the period the
    release reports — not a position. Columns under a "percent change" heading
    are excluded, and a release whose index column cannot be identified is
    rejected rather than approximated.
    """
    if fp.first_data_row is None or fp.label_col is None:
        raise Rejected("no data block found")

    weight_col = _weight_column(grid, fp)
    data_row = grid[fp.first_data_row]
    numeric = [
        c
        for c in range(fp.label_col + 1, min(MAX_SCAN_COLS, len(data_row)))
        if c != weight_col and is_number(data_row[c])
    ]
    if not numeric:
        raise Rejected("no numeric value columns")

    index_cols = []
    for c in numeric:
        label = _group_label(grid, fp, c)
        is_change = any(w in label for w in CHANGE_WORDS)
        is_index = any(w in label for w in INDEX_WORDS)
        if is_change and not is_index:
            continue
        if label and not is_index:
            continue
        index_cols.append(c)
    if not index_cols:
        raise Rejected("no index columns (all under a percent-change heading)")

    current = None
    if want is not None:
        for c in index_cols:
            if column_period(grid, fp, c) == want:
                current = c
                break
        if current is None:
            seen = [column_period(grid, fp, c) for c in index_cols]
            raise Rejected(f"no index column for period {want}; headers say {seen}")
    else:
        current = index_cols[-1]

    alt_col = _other_language_column(grid, fp)

    out: list[Observation] = []
    blanks = 0
    for row in grid[fp.first_data_row :]:
        if not row or not is_number(row[0]):
            # Tables contain blank spacer rows between groups, so end the table
            # only after a run of them -- stopping at the first would truncate
            # the block and silently drop later categories.
            blanks += 1
            if out and blanks >= 4:
                break
            continue
        blanks = 0
        level_num = as_number(row[0])
        if level_num is None or not float(level_num).is_integer():
            continue
        label_cell = row[fp.label_col] if fp.label_col < len(row) else None
        label_en, label_ar = split_bilingual(label_cell)
        # Some layouts put the two languages in one cell, others in two separate
        # columns; fill whichever half is still missing from the other column.
        if alt_col is not None and alt_col < len(row):
            alt_en, alt_ar = split_bilingual(row[alt_col])
            label_en = label_en or alt_en
            label_ar = label_ar or alt_ar
        if not (label_en or label_ar):
            continue
        value = as_number(row[current]) if current < len(row) else None
        if value is None:
            continue
        out.append(
            Observation(
                level=int(level_num),
                label_en=label_en,
                label_ar=label_ar,
                weight=as_number(row[weight_col]) if weight_col and weight_col < len(row) else None,
                value=value,
                column=current,
            )
        )
    if not out:
        raise Rejected("data block parsed to zero rows")
    return out


RECIPES: list[Recipe] = [
    Recipe(
        name="indicator-table-v1",
        claims=lambda fp: (
            fp.has_level_col
            and fp.first_data_row is not None
            and fp.label_col is not None
            and fp.n_numeric_cols >= 2
        ),
        read=read_indicator_table,
        note="level / label / weight / period columns; reads the trailing period",
    ),
]


def parse_sheet(
    grid: Grid, n_sheets: int, want: tuple[int, int] | None = None
) -> tuple[str, list[Observation]]:
    """Fingerprint, pick a recipe, read. Raises Rejected if nothing claims it."""
    fp = fingerprint(grid, n_sheets)
    for recipe in RECIPES:
        if recipe.claims(fp):
            return recipe.name, recipe.read(grid, fp, want)
    raise Rejected(f"no recipe claims shape {fp.key()}")
